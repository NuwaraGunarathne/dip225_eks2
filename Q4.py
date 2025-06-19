from openpyxl import load_workbook
import math

file = load_workbook('sagatave_eksamenam.xlsx')
sheet = file['Lapa_0']

header_idx = 3
cols = [sheet.cell(row=header_idx, column=i).value for i in range(1, sheet.max_column + 1)]

item_col = cols.index('Produkts') + 1
cost_col = cols.index('Cena') + 1

sum_prices = 0
match_items = 0

for r in range(header_idx + 1, sheet.max_row + 1):
    item = sheet.cell(row=r, column=item_col).value
    cost = sheet.cell(row=r, column=cost_col).value

    if item and 'laserjet' in str(item).lower():
        try:
            cost_val = float(cost)
        except (TypeError, ValueError):
            continue
        sum_prices += cost_val
        match_items += 1

if match_items:
    avg_price = sum_prices / match_items
    print("Rounded down average price:", math.floor(avg_price))
else:
    print("No LaserJet items found.")
