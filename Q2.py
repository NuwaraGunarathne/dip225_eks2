from openpyxl import load_workbook

excel_file = load_workbook('sagatave_eksamenam.xlsx')
sheet_data = excel_file['Lapa_0']

high_priority_2015 = 0

for i in range(2, sheet_data.max_row + 1):
    prio = sheet_data['H' + str(i)].value 
    date = sheet_data['J' + str(i)].value  

 
    if prio == 'High' and date and hasattr(date, 'year') and date.year == 2015:
        high_priority_2015 += 1

print("Total rows with High priority and 2015 date:", high_priority_2015)
