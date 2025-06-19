from openpyxl import load_workbook

workbook = load_workbook('sagatave_eksamenam.xlsx')
sheet = workbook['Lapa_0']

match_count = 0


for idx in range(2, sheet.max_row + 1):
    street = sheet['D' + str(idx)].value 
    town = sheet['E' + str(idx)].value   

    if street == "Adulienas iela" and town in ("Valmiera", "Saulkrasti"):
        match_count += 1

print("Total matching entries found:", match_count)
