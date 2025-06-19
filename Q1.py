from openpyxl import load_workbook

workbook = load_workbook('sagatave_eksamenam.xlsx')
sheet = workbook['Lapa_0']

header_index = 3
column_names = [sheet.cell(row=header_index, column=col).value for col in range(1, sheet.max_column + 1)]
address_index = column_names.index('Adrese') + 1
quantity_index = column_names.index('Skaits') + 1


match_counter = 0

for r in range(header_index + 1, sheet.max_row + 1):
    addr = sheet.cell(row=r, column=address_index).value
    qty = sheet.cell(row=r, column=quantity_index).value
    try:
        qty_val = float(qty)
    except (TypeError, ValueError):
        continue
    if addr and str(addr).startswith('Ain') and qty_val < 40:
        match_counter += 1

print("Total entries found:", match_counter)
